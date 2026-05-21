r"""Parse FTS XLSX rows and resolve Norwegian orgnrs via VAT.

FTS ``VAT number of beneficiary`` carries ``NO\d{9}MVA`` for most
Norwegian entities.  Some have ``-`` or ``*****`` (masked).
Country column uses full name ``Norway``.

Column indices (0-based, stable across 2007-2024):
    [0] Year, [4] Name, [5] VAT, [12] Country, [16] contracted_amount,
    [25] Subject, [26] Department, [29] Programme, [30] Funding type
"""

import re
import hashlib
import io

VAT_RE = re.compile(r"^NO(\d{9})MVA$")

COL_YEAR = 0
COL_NAME = 4
COL_VAT = 5
COL_COORDINATOR = 8
COL_COUNTRY = 12
COL_CONTRACTED = 16
COL_SUBJECT = 25
COL_DEPARTMENT = 26
COL_PROGRAMME = 29
COL_FUNDING_TYPE = 30
COL_START = 33
COL_END = 34
COL_CONTRACT_TYPE = 35
COL_MGMT_TYPE = 36


def resolve_orgnr_from_vat(vat_str):
    if not vat_str or vat_str in ("*****", "-", ""):
        return None, None
    m = VAT_RE.match(str(vat_str).strip())
    if m:
        return m.group(1), "vat"
    return None, None


def content_hash(row_dict):
    tracked = [
        str(row_dict.get("contracted_amount") or ""),
        str(row_dict.get("programme") or ""),
        str(row_dict.get("department") or ""),
        str(row_dict.get("subject") or ""),
        str(row_dict.get("funding_type") or ""),
    ]
    return hashlib.sha256("|".join(tracked).encode()).hexdigest()[:16]


def parse_xlsx(xlsx_bytes, fts_year):
    """Parse an FTS XLSX file and extract Norwegian rows.

    Args:
        xlsx_bytes: Raw XLSX bytes.
        fts_year: Integer year for metadata.

    Returns:
        Tuple of (resolved_rows, unresolved_rows).
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    resolved = []
    unresolved = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        country = str(row[COL_COUNTRY]) if row[COL_COUNTRY] else ""
        if "Norway" not in country:
            continue

        vat = str(row[COL_VAT]) if row[COL_VAT] else ""
        orgnr, method = resolve_orgnr_from_vat(vat)

        entry = {
            "orgnr": orgnr,
            "orgnr_resolution_method": method,
            "fts_year": fts_year,
            "name": str(row[COL_NAME]) if row[COL_NAME] else "",
            "vat": vat,
            "country": country,
            "contracted_amount": row[COL_CONTRACTED],
            "subject": str(row[COL_SUBJECT]) if row[COL_SUBJECT] else "",
            "department": str(row[COL_DEPARTMENT]) if row[COL_DEPARTMENT] else "",
            "programme": str(row[COL_PROGRAMME]) if row[COL_PROGRAMME] else "",
            "funding_type": str(row[COL_FUNDING_TYPE]) if row[COL_FUNDING_TYPE] else "",
            "coordinator": str(row[COL_COORDINATOR]) if row[COL_COORDINATOR] else "",
            "start_date": str(row[COL_START]) if row[COL_START] else "",
            "end_date": str(row[COL_END]) if row[COL_END] else "",
            "contract_type": str(row[COL_CONTRACT_TYPE]) if row[COL_CONTRACT_TYPE] else "",
            "management_type": str(row[COL_MGMT_TYPE]) if row[COL_MGMT_TYPE] else "",
        }
        entry["content_hash"] = content_hash(entry)

        if orgnr:
            resolved.append(entry)
        else:
            unresolved.append(entry)

    wb.close()
    return resolved, unresolved
